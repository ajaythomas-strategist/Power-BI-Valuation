"""
Professional Excel (.xlsx) exporter for Power BI Evaluation results.
Generates a multi-sheet formatted workbook using openpyxl:
Sheet 1: Summary
Sheet 2: Question-wise Evaluation
Sheet 3: Exceptions
Sheet 4: Evaluation Rules
"""
import io
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.schemas import (
    BatchEvaluationSummary,
    StudentEvaluationResult,
    EvaluationException,
    EvaluationRuleSet,
    EvaluationStatus,
)


class ExcelExporter:
    """Generates a styled .xlsx report matching institutional assessment standards."""

    # Brand Colors
    HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    
    REGULAR_FONT = Font(name="Segoe UI", size=10, color="0F172A")
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    # Status Fills & Fonts
    STATUS_STYLES = {
        "Correct": {
            "fill": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="166534"),
        },
        "Partially Correct": {
            "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="92400E"),
        },
        "Incorrect": {
            "fill": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="991B1B"),
        },
        "Exception": {
            "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="475569"),
        },
    }

    THIN_BORDER = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    TOTAL_BORDER = Border(
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="double", color="0F172A"),
    )

    @classmethod
    def generate_workbook(cls, summary: BatchEvaluationSummary) -> io.BytesIO:
        wb = Workbook()
        
        # 1. Summary Sheet
        ws_summary = wb.active
        ws_summary.title = "Summary"
        cls._build_summary_sheet(ws_summary, summary.results, summary.total_students, summary.rule_set)

        # 2. Question-wise Evaluation Sheet
        ws_details = wb.create_sheet(title="Question-wise Evaluation")
        cls._build_details_sheet(ws_details, summary.results)

        # 3. Exceptions Sheet
        ws_exceptions = wb.create_sheet(title="Exceptions")
        cls._build_exceptions_sheet(ws_exceptions, summary.exceptions)

        # 4. Evaluation Rules Sheet
        ws_rules = wb.create_sheet(title="Evaluation Rules")
        cls._build_rules_sheet(ws_rules, summary.rule_set)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def _build_summary_sheet(
        cls, ws, results: List[StudentEvaluationResult], total_students: int, rule_set: EvaluationRuleSet
    ):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Register No.",
            "Total Marks",
            "Maximum Marks",
            "Percentage",
            "Questions Evaluated",
            "Questions Correct",
            "Questions Partially Correct",
            "Questions Incorrect",
            "Evaluation Status",
        ]

        # Write Headers
        ws.append(headers)
        cls._format_header_row(ws, 1, len(headers))

        row_idx = 2
        for res in results:
            ws.append([
                res.register_no,
                res.total_marks,
                res.maximum_marks,
                res.percentage / 100.0,
                res.questions_evaluated,
                res.questions_correct,
                res.questions_partially_correct,
                res.questions_incorrect,
                res.status.value,
            ])

            # Apply cell styles
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = cls.REGULAR_FONT
                cell.border = cls.THIN_BORDER
                cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx in (1, 5, 6, 7, 8, 9) else "right")

            # Format Percentage
            pct_cell = ws.cell(row=row_idx, column=4)
            pct_cell.number_format = "0.0%"

            # Format Status Tag
            status_cell = ws.cell(row=row_idx, column=9)
            st_style = cls.STATUS_STYLES.get(res.status.value, cls.STATUS_STYLES["Exception"])
            status_cell.fill = st_style["fill"]
            status_cell.font = st_style["font"]

            row_idx += 1

        # Summary Average Row
        if len(results) > 0:
            avg_row = [
                "Average / Summary",
                f"=AVERAGE(B2:B{row_idx-1})",
                rule_set.total_marks,
                f"=AVERAGE(D2:D{row_idx-1})",
                f"=AVERAGE(E2:E{row_idx-1})",
                f"=SUM(F2:F{row_idx-1})",
                f"=SUM(G2:G{row_idx-1})",
                f"=SUM(H2:H{row_idx-1})",
                f"{len(results)} Evaluated",
            ]
            ws.append(avg_row)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = cls.BOLD_FONT
                cell.border = cls.TOTAL_BORDER
                if col_idx == 4:
                    cell.number_format = "0.0%"
                elif col_idx in (2, 3, 5):
                    cell.number_format = "0.0"

        cls._autofit_columns(ws)

    @classmethod
    def _build_details_sheet(cls, ws, results: List[StudentEvaluationResult]):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Register No.",
            "Question ID",
            "Question",
            "Requirement",
            "Expected",
            "Student Implementation",
            "Result",
            "Marks Awarded",
            "Maximum Marks",
            "Remarks",
        ]

        ws.append(headers)
        cls._format_header_row(ws, 1, len(headers))

        row_idx = 2
        for res in results:
            for q in res.question_details:
                ws.append([
                    res.register_no,
                    q.question_id,
                    q.question_text,
                    q.requirement_summary,
                    q.expected_summary,
                    q.student_summary,
                    q.status.value,
                    q.marks_awarded,
                    q.maximum_marks,
                    q.remarks,
                ])

                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = cls.REGULAR_FONT
                    cell.border = cls.THIN_BORDER
                    align = Alignment(vertical="center", wrap_text=(col_idx in (3, 4, 5, 6, 10)))
                    if col_idx in (1, 2, 7):
                        align.horizontal = "center"
                    elif col_idx in (8, 9):
                        align.horizontal = "right"
                    cell.alignment = align

                # Status color
                status_cell = ws.cell(row=row_idx, column=7)
                st_style = cls.STATUS_STYLES.get(q.status.value, cls.STATUS_STYLES["Incorrect"])
                status_cell.fill = st_style["fill"]
                status_cell.font = st_style["font"]

                row_idx += 1

        cls._autofit_columns(ws, max_len=45)

    @classmethod
    def _build_exceptions_sheet(cls, ws, exceptions: List[EvaluationException]):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Register No.",
            "Issue Type",
            "Description",
            "Status",
            "Diagnostic Details",
        ]

        ws.append(headers)
        cls._format_header_row(ws, 1, len(headers))

        row_idx = 2
        if not exceptions:
            ws.append(["None", "No exceptions encountered", "All submissions processed successfully.", "OK", ""])
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=2, column=col_idx)
                cell.font = cls.REGULAR_FONT
                cell.border = cls.THIN_BORDER
        else:
            for exc in exceptions:
                ws.append([
                    exc.register_no,
                    exc.issue_type.value if hasattr(exc.issue_type, "value") else str(exc.issue_type),
                    exc.description,
                    exc.status,
                    exc.details or "",
                ])

                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = cls.REGULAR_FONT
                    cell.border = cls.THIN_BORDER
                    cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx in (1, 4) else "left")

                # Exception status style
                status_cell = ws.cell(row=row_idx, column=4)
                status_cell.fill = cls.STATUS_STYLES["Exception"]["fill"]
                status_cell.font = cls.STATUS_STYLES["Exception"]["font"]

                row_idx += 1

        cls._autofit_columns(ws)

    @classmethod
    def _build_rules_sheet(cls, ws, rule_set: EvaluationRuleSet):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Question ID",
            "Question Text",
            "Criterion ID",
            "Criterion Title",
            "Target Component",
            "Target Property",
            "Expected Value",
            "Accepted Variations",
            "Allocated Marks",
        ]

        ws.append(headers)
        cls._format_header_row(ws, 1, len(headers))

        row_idx = 2
        for q in rule_set.questions:
            for c in q.criteria:
                variations_str = ", ".join(c.accepted_variations + q.accepted_variations)
                ws.append([
                    q.question_id,
                    q.question_text,
                    c.id,
                    c.title,
                    c.component.value if hasattr(c.component, "value") else str(c.component),
                    c.target_property,
                    str(c.expected_value),
                    variations_str or "None",
                    c.marks,
                ])

                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = cls.REGULAR_FONT
                    cell.border = cls.THIN_BORDER
                    cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx in (1, 3, 5, 6) else ("right" if col_idx == 9 else "left"))

                row_idx += 1

        cls._autofit_columns(ws)

    @classmethod
    def _format_header_row(cls, ws, row_idx: int, num_cols: int):
        ws.row_dimensions[row_idx].height = 28
        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cls.THIN_BORDER

    @classmethod
    def _autofit_columns(cls, ws, max_len: int = 50):
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_width = 0
            for cell in col:
                val = str(cell.value or "")
                # Ignore formula length
                if val.startswith("="):
                    val = "100.0%"
                max_width = max(max_width, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_width + 4, 12), max_len)
